from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "bbdd.md"
OUTPUT = ROOT / "Base de datos actualizada - La Galana.docx"

ACCENT = "6F7F45"
ACCENT_DARK = "354024"
SOFT = "F4F0E7"
TABLE_HEADER = "E6DCC7"
TEXT = "1F241A"
MUTED = "5F655A"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, bold=False, italic=False, color=TEXT, size=10.5, font="Aptos"):
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def add_rich_text(paragraph, text, size=10.5, color=TEXT, bold_default=False):
    parts = text.split("`")
    for index, part in enumerate(parts):
        if not part:
            continue
        is_code = index % 2 == 1
        run = paragraph.add_run(part)
        style_run(
            run,
            bold=bold_default,
            color=color if not is_code else ACCENT_DARK,
            size=size,
            font="Cascadia Mono" if is_code else "Aptos",
        )


def add_horizontal_rule(paragraph):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), ACCENT)
    border.append(bottom)
    p_pr.append(border)


def apply_table_style(table):
    table.style = "Table Grid"
    table.autofit = True

    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            set_repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, TABLE_HEADER)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    style_run(
                        run,
                        bold=row_index == 0,
                        color=ACCENT_DARK if row_index == 0 else TEXT,
                        size=9 if row_index == 0 else 8.5,
                    )


def add_table(document, rows):
    if len(rows) < 2:
        return

    column_count = len(rows[0])
    table = document.add_table(rows=1, cols=column_count)
    header = table.rows[0].cells
    for index, value in enumerate(rows[0]):
        header[index].text = clean_markdown_inline(value)

    for row_values in rows[1:]:
        cells = table.add_row().cells
        for index in range(column_count):
            value = row_values[index] if index < len(row_values) else ""
            cells[index].text = clean_markdown_inline(value)

    apply_table_style(table)
    document.add_paragraph()


def clean_markdown_inline(text):
    return text.replace("`", "").replace("**", "").strip()


def is_separator_row(line):
    stripped = line.strip()
    return stripped.startswith("|") and set(stripped.replace("|", "").replace("-", "").replace(":", "").strip()) == set()


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        if not is_separator_row(lines[index]):
            cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
            rows.append(cells)
        index += 1
    return rows, index


def add_code_block(document, code_lines):
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F7F3")
    set_cell_margins(cell, top=130, start=160, bottom=130, end=160)

    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for line_index, line in enumerate(code_lines):
        if line_index > 0:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        style_run(run, color="263021", size=8.2, font="Cascadia Mono")

    document.add_paragraph()


def add_section_label(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text.upper())
    style_run(run, bold=True, color=ACCENT, size=8.5)


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)

    for name, size, color in [
        ("Heading 1", 17, ACCENT_DARK),
        ("Heading 2", 13, ACCENT_DARK),
        ("Heading 3", 11, ACCENT),
    ]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run("Casa Rural La Galana")
    style_run(title_run, bold=True, color=ACCENT, size=12)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle_run = subtitle.add_run("Base de datos actualizada")
    style_run(subtitle_run, bold=True, color=ACCENT_DARK, size=24, font="Aptos Display")
    add_horizontal_rule(subtitle)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich_text(meta, "Apartado técnico para la memoria del proyecto intermodular", size=10, color=MUTED)
    meta.paragraph_format.space_after = Pt(14)

    lines = SOURCE.read_text(encoding="utf-8-sig").splitlines()
    index = 0
    in_code = False
    code_lines = []

    while index < len(lines):
        line = lines[index].rstrip()

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not line.strip():
            index += 1
            continue

        if line.strip().startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue

        if line.startswith("# "):
            index += 1
            continue

        if line.startswith("## "):
            heading = clean_markdown_inline(line[3:])
            if heading.lower().startswith("table ") or heading.lower().startswith("tabla "):
                add_section_label(document, "Tabla de base de datos")
            paragraph = document.add_heading(heading, level=1)
            paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue

        if line.startswith("### "):
            paragraph = document.add_heading(clean_markdown_inline(line[4:]), level=2)
            paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue

        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_rich_text(paragraph, line[2:], size=10)
            index += 1
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.12
        add_rich_text(paragraph, clean_markdown_inline(line) if line.startswith("|") else line)
        index += 1

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("La Galana Manager · Documentacion de base de datos")
    style_run(footer_run, color=MUTED, size=8.5)

    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
