import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "_tfg_pdf_extracted.json"
OUTPUT = ROOT / "TFG 3º Entrega.docx"

ACCENT = "6F7F45"
ACCENT_DARK = "2F3B24"
TABLE_HEADER = "E6DCC7"
TEXT = "20231D"
MUTED = "62685D"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=100, bottom=85, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, size=10.5, bold=False, italic=False, color=TEXT, font="Aptos"):
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def is_heading(line):
    if re.match(r"^\d+(\.\d+)*\.\s+", line):
        return True
    if re.match(r"^\d+\.\s+", line):
        return True
    if line in {
        "Contenido",
        "Modelo relacional revisado",
        "Tablas de base de datos detalladas",
        "Estructura de clases/componentes actual",
        "Lógica actual de la aplicación",
        "Descripción textual de los casos de uso",
        "Batería de pruebas detallada",
        "Estado actual de la aplicación por bloques",
        "Anexo I. Estructura principal del proyecto",
        "Anexo II. Variables de entorno",
        "Anexo III. Verificación ejecutada",
    }:
        return True
    return False


def heading_level(line):
    if re.match(r"^\d+\.\s+", line):
        return 1
    if re.match(r"^\d+\.\d+\.\s+", line):
        return 2
    if re.match(r"^\d+\.\d+\.\d+\.\s+", line):
        return 3
    return 2


def looks_like_table(lines, start):
    if start + 2 >= len(lines):
        return False
    shortish = [len(lines[start + offset]) <= 55 for offset in range(min(5, len(lines) - start))]
    return sum(shortish) >= 4 and not is_heading(lines[start])


def collect_table(lines, start):
    rows = []
    index = start
    while index < len(lines):
        line = lines[index]
        if not line.strip() or line.startswith("•") or is_heading(line):
            break
        if len(line) > 80 and rows:
            break
        rows.append([line])
        index += 1
        if len(rows) >= 16:
            break
    return rows, index


def add_simple_table(document, rows):
    if len(rows) < 2:
        return False
    table = document.add_table(rows=0, cols=1)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        cells[0].text = row_values[0]
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[0])
        if row_index == 0:
            set_cell_shading(cells[0], TABLE_HEADER)
            set_repeat_table_header(table.rows[0])
        for paragraph in cells[0].paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                style_run(run, size=9.2, bold=row_index == 0, color=ACCENT_DARK if row_index == 0 else TEXT)
    document.add_paragraph()
    return True


def add_paragraph(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    style_run(run, size=10.5)
    return paragraph


def add_code_like(document, lines):
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F7F3")
    set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    paragraph = cell.paragraphs[0]
    for i, line in enumerate(lines):
        if i:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        style_run(run, size=8.5, font="Cascadia Mono", color=ACCENT_DARK)
    document.add_paragraph()


def add_image(document, image_path, page_number, index):
    path = Path(image_path)
    if not path.exists():
        return
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(f"Figura extraida de la pagina {page_number}")
    style_run(run, size=9, italic=True, color=MUTED)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        paragraph.add_run().add_picture(str(path), width=Cm(14.5))
    except Exception:
        fallback = paragraph.add_run(f"[Imagen no insertada: {path.name}]")
        style_run(fallback, size=9, italic=True, color=MUTED)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [
        ("Heading 1", 17, ACCENT_DARK),
        ("Heading 2", 13, ACCENT_DARK),
        ("Heading 3", 11.5, ACCENT),
    ]:
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)


def main():
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    document = Document()
    configure_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("LA GALANA MANAGER")
    style_run(run, size=24, bold=True, color=ACCENT_DARK, font="Aptos Display")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Memoria del proyecto intermodular")
    style_run(run, size=12, color=MUTED)
    document.add_paragraph()

    for page in data["pages"]:
        lines = [line.strip() for line in page["text"].splitlines() if line.strip()]
        filtered = [line for line in lines if not line.startswith("Memoria - Casa Rural La Galana - Página")]
        index = 0

        while index < len(filtered):
            line = filtered[index]
            if line.startswith("LA GALANA MANAGER") and page["number"] == 1:
                index += 1
                continue

            if line.startswith("•"):
                paragraph = document.add_paragraph(style="List Bullet")
                run = paragraph.add_run(line[1:].strip())
                style_run(run, size=10.3)
                index += 1
                continue

            if line.startswith(("const ", "export ", ".from(", ".select(", ".lte(", ".gte(")):
                code = []
                while index < len(filtered) and (
                    filtered[index].startswith(("const ", "export ", ".from(", ".select(", ".lte(", ".gte("))
                    )
                    or filtered[index].strip() in {";", "};"}
                ):
                    code.append(filtered[index])
                    index += 1
                add_code_like(document, code)
                continue

            if is_heading(line):
                paragraph = document.add_heading(line, level=heading_level(line))
                paragraph.paragraph_format.keep_with_next = True
                index += 1
                continue

            add_paragraph(document, line)
            index += 1

        for image_index, image_path in enumerate(page.get("images", []), start=1):
            add_image(document, image_path, page["number"], image_index)

        if page["number"] < data["page_count"]:
            document.add_section(WD_SECTION.NEW_PAGE)

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("La Galana Manager · Memoria convertida desde PDF")
    style_run(footer_run, size=8.5, color=MUTED)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
